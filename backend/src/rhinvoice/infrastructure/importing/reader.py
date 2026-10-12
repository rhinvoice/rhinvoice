"""Read CSV, XLSX and XLS files into a table of texts.

The type comes from the content, not the file name: XLSX files are ZIP archives,
XLS files (as older invoicing programs export them) OLE compound documents, anything else is
read as CSV. CSV files from Windows programs are often Windows-1252 and separated
by semicolons; both are detected.
"""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import datetime, time
from typing import TYPE_CHECKING

import xlrd
from openpyxl import load_workbook

from rhinvoice.application.importing.ports import ImportTable
from rhinvoice.domain.shared.errors import ValidationError

if TYPE_CHECKING:
    from collections.abc import Iterable

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_SCANNED_ROWS = 5 * MAX_ROWS
"""Empty rows count too: a spreadsheet may declare a million of them."""
MAX_UNPACKED_BYTES = 50 * 1024 * 1024
"""An XLSX file is a ZIP archive; a few MB can unpack to gigabytes."""
_XLSX_SIGNATURE = b"PK\x03\x04"
_XLS_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_DELIMITERS = ";,\t|"
_SNIFF_BYTES = 8192


def _unreadable(message: str) -> ValidationError:
    return ValidationError("file", message, code="unreadable_file")


def _text(value: object) -> str:
    """Write a cell as a person would type it: 12 not 12.0, dates as 2026-03-09."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float):
        # Binary fractions: 1234.5000000001 is what the person typed as 1234.5.
        rounded = round(value, 9)
        return str(int(rounded)) if rounded.is_integer() else repr(rounded)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    return str(value).strip()


def _too_many_rows() -> ValidationError:
    return ValidationError("file", f"the file has more than {MAX_ROWS} rows", code="too_many_rows")


def _table(rows: Iterable[list[str]], **details: str | None) -> ImportTable:
    """Header row first; empty rows dropped, every row as wide as the widest.

    Each cell is looked at once, and reading stops at the first row beyond the
    limits: a crafted file must not keep the server busy.
    """
    cleaned: list[list[str]] = []
    width = 0
    for scanned, row in enumerate(rows, start=1):
        if scanned > MAX_SCANNED_ROWS:
            raise _too_many_rows()
        # PostgreSQL texts cannot hold NUL characters.
        cells = [cell.replace("\x00", "").strip() for cell in row]
        # Trailing empty cells (e.g. formatted but empty ones) do not count.
        while cells and not cells[-1]:
            cells.pop()
        if len(cells) > MAX_COLUMNS:
            raise ValidationError(
                "file", f"the file has more than {MAX_COLUMNS} columns", code="too_many_columns"
            )
        if cells:
            cleaned.append(cells)
            width = max(width, len(cells))
            if len(cleaned) > MAX_ROWS + 1:
                raise _too_many_rows()
    if not cleaned:
        raise ValidationError("file", "the file is empty", code="empty_file")
    padded = [tuple(row + [""] * (width - len(row))) for row in cleaned]
    return ImportTable(
        # Columns without a name keep an empty one; the page names them by position.
        headers=padded[0],
        rows=tuple(padded[1:]),
        file_type=details["file_type"] or "csv",
        encoding=details.get("encoding"),
        delimiter=details.get("delimiter"),
    )


def _decode(content: bytes) -> tuple[str, str]:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        # Excel's "Unicode text" export.
        return content.decode("utf-16", errors="replace"), "UTF-16"
    if content.startswith(b"\xef\xbb\xbf"):
        return content[3:].decode("utf-8", errors="replace"), "UTF-8"
    try:
        return content.decode("utf-8"), "UTF-8"
    except UnicodeDecodeError:
        # Windows-1252 leaves five bytes undefined; those become replacement characters.
        return content.decode("cp1252", errors="replace"), "Windows-1252"


def _delimiter(text: str) -> str:
    sample = text[:_SNIFF_BYTES]
    try:
        return csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        # One column, or too irregular for the sniffer: the most frequent candidate.
        first_line = sample.splitlines()[0] if sample else ""
        counts = {candidate: first_line.count(candidate) for candidate in _DELIMITERS}
        best = max(counts, key=lambda candidate: counts[candidate])
        return best if counts[best] else ";"


def _read_csv(content: bytes) -> ImportTable:
    text, encoding = _decode(content)
    delimiter = _delimiter(text)
    try:
        rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    except csv.Error as error:
        raise _unreadable(f"the CSV file cannot be read: {error}") from error
    return _table(rows, file_type="csv", encoding=encoding, delimiter=delimiter)


def _unpacked_size(content: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return sum(info.file_size for info in archive.infolist())


def _read_xlsx(content: bytes) -> ImportTable:
    try:
        unpacked = _unpacked_size(content)
    except zipfile.BadZipFile as error:
        raise _unreadable("the Excel file (.xlsx) cannot be read") from error
    if unpacked > MAX_UNPACKED_BYTES:
        raise ValidationError(
            "file", "the Excel file unpacks to more than 50 MB", code="file_too_large"
        )
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    # Damaged or crafted files make the parsers raise all kinds of errors.
    except Exception as error:
        raise _unreadable("the Excel file (.xlsx) cannot be read") from error
    try:
        sheet = workbook.worksheets[0]
        # The size a file declares ("A1:XFD1048576") may be made up: rows are read
        # as long as they actually are.
        sheet.reset_dimensions()
        rows = ([_text(value) for value in row] for row in sheet.iter_rows(values_only=True))
        return _table(rows, file_type="xlsx")
    finally:
        workbook.close()


def _read_xls(content: bytes) -> ImportTable:
    try:
        book = xlrd.open_workbook(file_contents=content)
    # Damaged files, see _read_xlsx.
    except Exception as error:
        raise _unreadable("the Excel file (.xls) cannot be read") from error
    sheet = book.sheet_by_index(0)

    def cell(row: int, column: int) -> str:
        item = sheet.cell(row, column)
        if item.ctype == xlrd.XL_CELL_DATE and isinstance(item.value, float):
            return _text(xlrd.xldate.xldate_as_datetime(item.value, book.datemode))
        return _text(item.value)

    rows = ([cell(row, column) for column in range(sheet.ncols)] for row in range(sheet.nrows))
    return _table(rows, file_type="xls")


class SpreadsheetReader:
    def read(self, content: bytes) -> ImportTable:
        if len(content) > MAX_FILE_BYTES:
            raise ValidationError("file", "the file is larger than 5 MB", code="file_too_large")
        if content.startswith(_XLSX_SIGNATURE):
            return _read_xlsx(content)
        if content.startswith(_XLS_SIGNATURE):
            return _read_xls(content)
        return _read_csv(content)
