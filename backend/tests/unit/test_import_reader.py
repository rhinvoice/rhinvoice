"""Reading CSV, XLSX and XLS files for the import."""

import io
import time
import zipfile
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from rhinvoice.domain.shared.errors import ValidationError
from rhinvoice.infrastructure.importing import reader as reader_module
from rhinvoice.infrastructure.importing.reader import (
    MAX_COLUMNS,
    MAX_FILE_BYTES,
    MAX_ROWS,
    MAX_SCANNED_ROWS,
    SpreadsheetReader,
)
from tests.helpers import raises_exactly

FIXTURES = Path(__file__).parent.parent / "fixtures"
read = SpreadsheetReader().read
# Excel has no time zones: its date-times are naive.
MORNING = datetime(2026, 3, 9, 8, 30)  # noqa: DTZ001 - see above
MIDNIGHT = datetime(2026, 3, 10)  # noqa: DTZ001 - see above


def test_csv_from_windows_programs() -> None:
    content = "KdNr;Firma;Ort\r\n10001;Bergbahnen Müster AG;Buchs SG\r\n\r\n;;\r\n".encode("cp1252")
    table = read(content)
    assert table.headers == ("KdNr", "Firma", "Ort")
    assert table.rows == (("10001", "Bergbahnen Müster AG", "Buchs SG"),)
    assert (table.file_type, table.encoding, table.delimiter) == ("csv", "Windows-1252", ";")


def test_utf8_csv_with_bom_quotes_and_commas() -> None:
    content = '﻿Name,Ort,Notiz\n"Muster, Anna",Zürich,"Zeile 1\nZeile 2"\n'.encode()
    table = read(content)
    assert table.headers == ("Name", "Ort", "Notiz")
    assert table.rows == (("Muster, Anna", "Zürich", "Zeile 1\nZeile 2"),)
    assert (table.encoding, table.delimiter) == ("UTF-8", ",")
    assert read("Name;Ort\nAnna;Zürich\n".encode()).encoding == "UTF-8"


def test_tabs_single_columns_and_ragged_rows() -> None:
    assert read(b"Name\tOrt\nAnna\tBern\n").delimiter == "\t"
    single = read(b"Name\nAnna\nBeat\n")
    assert (single.headers, single.rows, single.delimiter) == (
        ("Name",),
        (("Anna",), ("Beat",)),
        ";",
    )
    ragged = read(b"Name;;Ort\nAnna;x;Bern;extra\nBeat\n")
    # Columns without a name stay without one; the page names them by position.
    assert ragged.headers == ("Name", "", "Ort", "")
    assert ragged.rows == (("Anna", "x", "Bern", "extra"), ("Beat", "", "", ""))
    # A one-line file whose separator the sniffer cannot tell.
    assert read(b"Name|Ort").delimiter == "|"


def test_xlsx_cells_become_texts() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Nr", "Name", "Preis", "Datum", "Zeit", "Aktiv", None, None])
    sheet.append([12, " Anna ", 12.5, date(2026, 3, 9), MORNING, True, None, None])
    sheet.append([None, None, None, None, None, None, None, None])
    sheet.append([13.0, "Beat", 3, MIDNIGHT, None, False, None, None])
    sheet.append([14, "Carla", 1234.5000000001, None, None, None, None, None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    table = read(buffer.getvalue())
    assert table.file_type == "xlsx"
    assert table.headers == ("Nr", "Name", "Preis", "Datum", "Zeit", "Aktiv")
    assert table.rows == (
        ("12", "Anna", "12.5", "2026-03-09", "2026-03-09T08:30:00", "1"),
        ("13", "Beat", "3", "2026-03-10", "", "0"),
        # Binary fractions are rounded to what was typed.
        ("14", "Carla", "1234.5", "", "", ""),
    )
    assert (table.encoding, table.delimiter) == (None, None)


def test_old_format_xls() -> None:
    table = read((FIXTURES / "customers.xls").read_bytes())
    assert table.file_type == "xls"
    assert table.headers == (
        "Kundennummer",
        "Firma",
        "Straße",
        "PLZ",
        "Ort",
        "Rabatt",
        "Kunde seit",
    )
    assert table.rows[0] == (
        "10001",
        "Bergbahnen Müster AG",
        "Dorfstrasse 1",
        "9470",
        "Buchs SG",
        "2.5",
        "2014-03-01",
    )
    assert len(table.rows) == 2


def test_files_that_cannot_be_read() -> None:
    with raises_exactly(ValidationError, "the file is empty") as empty:
        read(b"\n;;\n")
    assert empty.value.code == "empty_file"
    with raises_exactly(ValidationError, "the file is larger than 5 MB") as large:
        read(b"x" * (MAX_FILE_BYTES + 1))
    assert large.value.code == "file_too_large"
    with raises_exactly(ValidationError, "the Excel file (.xlsx) cannot be read") as xlsx:
        read(b"PK\x03\x04 not really a zip")
    assert (xlsx.value.field, xlsx.value.code) == ("file", "unreadable_file")
    with raises_exactly(ValidationError, "the Excel file (.xls) cannot be read"):
        read(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600)
    with raises_exactly(ValidationError, f"the file has more than {MAX_ROWS} rows") as rows:
        read(b"Name\n" + b"Anna\n" * (MAX_ROWS + 1))
    assert rows.value.code == "too_many_rows"
    assert len(read(b"Name\n" + b"Anna\n" * MAX_ROWS).rows) == MAX_ROWS
    wide = ";".join(f"C{index}" for index in range(MAX_COLUMNS + 1)).encode()
    with raises_exactly(
        ValidationError, f"the file has more than {MAX_COLUMNS} columns"
    ) as columns:
        read(wide)
    assert columns.value.code == "too_many_columns"
    with raises_exactly(
        ValidationError, "the CSV file cannot be read: field larger than field limit (131072)"
    ):
        read(b'Name;Ort\n"' + b"x" * 131_073 + b'";Bern\n')
    assert read(b"Name;Ort\nAn\x00na;Bern\n").rows == (("Anna", "Bern"),)


def test_undefined_windows_1252_bytes_are_replaced() -> None:
    # 0x81 is not defined in Windows-1252 (and not valid UTF-8).
    table = read(b"Name\nA\x81B\xe4\n")
    assert table.rows == (("A�Bä",),)
    assert table.encoding == "Windows-1252"


def test_unicode_text_from_excel() -> None:
    table = read("Name\tOrt\nMüller\tZürich\n".encode("utf-16"))
    assert (table.encoding, table.delimiter) == ("UTF-16", "\t")
    assert table.rows == (("Müller", "Zürich"),)


def test_crafted_files_are_refused_quickly() -> None:
    # One row with a huge number of separators, then many rows: every cell is read once.
    started = time.monotonic()
    wide = b"x" + b";" * 200_000 + b"\n" + b"x\n" * MAX_ROWS
    assert len(read(wide).rows) == MAX_ROWS
    content_beyond = b"Name\n" + b";" * MAX_COLUMNS + b"x\n"
    with raises_exactly(ValidationError, f"the file has more than {MAX_COLUMNS} columns"):
        read(content_beyond)
    assert time.monotonic() - started < 5
    # Empty rows count against a limit of their own.
    with raises_exactly(ValidationError, f"the file has more than {MAX_ROWS} rows"):
        read(b"Name\n" + b";\n" * MAX_SCANNED_ROWS)


def test_xlsx_with_too_many_columns_is_refused() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Name"])
    sheet.cell(row=2, column=MAX_COLUMNS + 50, value="x")
    buffer = io.BytesIO()
    workbook.save(buffer)
    with raises_exactly(ValidationError, f"the file has more than {MAX_COLUMNS} columns"):
        read(buffer.getvalue())


def test_xlsx_that_unpacks_to_too_much_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(reader_module, "MAX_UNPACKED_BYTES", 1_000)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", "x" * 2_000)
    with raises_exactly(ValidationError, "the Excel file unpacks to more than 50 MB") as caught:
        read(buffer.getvalue())
    assert caught.value.code == "file_too_large"


def test_zip_files_that_are_no_workbook_cannot_be_read() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "no workbook")
    with raises_exactly(ValidationError, "the Excel file (.xlsx) cannot be read"):
        read(buffer.getvalue())
